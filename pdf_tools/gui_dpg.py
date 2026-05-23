from __future__ import annotations

import io
import queue
import subprocess
import sys
import threading
import time
import webbrowser
from datetime import UTC, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Callable

import dearpygui.dearpygui as dpg

from .authorization import (
    AuthorizationResult,
    AuthorizationStore,
    DEFAULT_SERVER_URL,
    TRIAL_DAILY_LIMIT,
    TrialUsageStore,
    get_authorization_status,
    get_device_fingerprint,
    import_offline_license,
    parse_iso_datetime,
    poll_device_binding,
    save_offline_device_request,
    start_device_binding,
)
from .converter import convert_pdf_to_docx
from .pages import pdf_compress, pdf_encrypt, pdf_merge, pdf_split, pdf_to_excel, pdf_to_image, pdf_to_ppt as pdf_to_ppt_page, pdf_to_word


APP_TITLE = "PDF Tools"
VIEWPORT_WIDTH = 1180
VIEWPORT_HEIGHT = 760
NATIVE_DIALOG_UNAVAILABLE = object()
TASK_BUTTON_TAGS = ["word_run", "merge_run", "split_run", "compress_run", "ppt_run", "image_run", "encrypt_run"]
APP_ICON_PATH = Path(__file__).resolve().parent.parent / "assets" / "icon" / "app.ico"
OFFLINE_DEVICE_REQUEST_FILENAME = "设备请求.json"
TOOLS = [
    ("merge", "PDF 合并", "合并多个 PDF 文件"),
    ("split", "PDF 分割", "按页数拆分 PDF"),
    ("compress", "PDF 压缩", "减小 PDF 文件体积"),
    ("word", "PDF 转 Word", "转换为 DOCX 文档"),
    ("ppt", "PDF 转 PPT", "每页生成一张幻灯片"),
    ("excel", "PDF 转 Excel", "表格提取功能"),
    ("image", "PDF 转图片", "导出 PNG/JPG 图片"),
    ("encrypt", "PDF 加密", "添加打开密码"),
]
TOOL_LABELS = {key: label for key, label, _description in TOOLS}


def format_local_datetime(value: str | None) -> str:
    if not value:
        return "未提供"
    try:
        return parse_iso_datetime(value).astimezone().strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return value


def format_remaining(value: str | None) -> str:
    if not value:
        return "未知"
    try:
        expires_at = parse_iso_datetime(value)
    except ValueError:
        return "未知"
    seconds = int((expires_at - datetime.now(UTC)).total_seconds())
    if seconds <= 0:
        return "已过期"
    days, remainder = divmod(seconds, 24 * 60 * 60)
    hours = remainder // (60 * 60)
    if days > 0:
        return f"剩余 {days} 天 {hours} 小时"
    minutes = (remainder % (60 * 60)) // 60
    return f"剩余 {hours} 小时 {minutes} 分钟"


def get_authorization_summary(result: AuthorizationResult) -> dict[str, str]:
    payload: dict[str, Any] = result.payload or {}
    source = {"online": "在线订阅", "offline": "离线 License"}.get(result.source, "未授权")
    expires_at = None
    edition = "-"

    if result.source == "offline":
        expires_at = str(payload.get("expires_at") or "") or None
        edition = str(payload.get("edition") or "-")
    elif result.source == "online":
        expires_at = str(payload.get("subscriptionExpiresAt") or payload.get("cacheUntil") or "") or None
        edition = str(payload.get("edition") or "-")

    return {
        "status": "已生效" if result.valid else "未授权",
        "source": source,
        "edition": edition,
        "expires_at": format_local_datetime(expires_at),
        "remaining": format_remaining(expires_at) if result.valid else "-",
        "message": result.message,
    }


class PdfToolsApp:
    def __init__(self) -> None:
        self.current_tool = "word"
        self.busy = False
        self.output_path: str | None = None
        self.auth_store = AuthorizationStore()
        self.trial_usage_store = TrialUsageStore()
        self.poll_code: str | None = None
        self.confirm_url: str | None = None
        self.last_poll_at = 0.0
        self.title_font = None
        self.word_pdf: str | None = None
        self.split_pdf: str | None = None
        self.compress_pdf: str | None = None
        self.ppt_pdf: str | None = None
        self.image_pdf: str | None = None
        self.encrypt_pdf: str | None = None
        self.merge_files: list[str] = []
        self._file_dialog_target: str | None = None
        self._results: queue.Queue[tuple[str, str | None]] = queue.Queue()

    def run(self) -> int:
        dpg.create_context()
        self._setup_fonts()
        self._setup_theme()
        self._build_ui()
        icon = str(APP_ICON_PATH) if APP_ICON_PATH.exists() else ""
        dpg.create_viewport(title=APP_TITLE, width=VIEWPORT_WIDTH, height=VIEWPORT_HEIGHT, small_icon=icon, large_icon=icon)
        dpg.setup_dearpygui()
        dpg.show_viewport()
        dpg.set_primary_window("main_window", True)
        dpg.start_dearpygui()
        dpg.destroy_context()
        return 0

    def _setup_fonts(self) -> None:
        candidates = [
            Path("C:/Windows/Fonts/msyh.ttc"),
            Path("C:/Windows/Fonts/simhei.ttf"),
            Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
            Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
        ]
        font_path = next((path for path in candidates if path.exists()), None)
        if not font_path:
            return
        try:
            with dpg.font_registry():
                with dpg.font(str(font_path), 18) as default_font:
                    pass
                with dpg.font(str(font_path), 22) as title_font:
                    pass
            dpg.bind_font(default_font)
            self.title_font = title_font
        except Exception:
            self.title_font = None

    def _setup_theme(self) -> None:
        with dpg.theme() as theme:
            with dpg.theme_component(dpg.mvAll):
                dpg.add_theme_color(dpg.mvThemeCol_WindowBg, (246, 248, 251), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_ChildBg, (255, 255, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_PopupBg, (255, 255, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Text, (31, 41, 55), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Button, (241, 245, 249), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_ButtonHovered, (226, 232, 240), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_ButtonActive, (203, 213, 225), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_FrameBg, (255, 255, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_FrameBgHovered, (243, 246, 250), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_FrameBgActive, (239, 246, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_InputTextCursor, (15, 23, 42), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_TextSelectedBg, (191, 219, 254), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_NavCursor, (37, 99, 235), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Border, (214, 222, 235), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_MenuBarBg, (255, 255, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Header, (239, 246, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_HeaderHovered, (219, 234, 254), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_HeaderActive, (191, 219, 254), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_TitleBg, (248, 250, 252), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_TitleBgActive, (248, 250, 252), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_TitleBgCollapsed, (248, 250, 252), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_ScrollbarBg, (248, 250, 252), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_ScrollbarGrab, (203, 213, 225), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_ScrollbarGrabHovered, (148, 163, 184), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_ScrollbarGrabActive, (100, 116, 139), category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_FrameRounding, 6, category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_ChildRounding, 8, category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_WindowRounding, 8, category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_WindowPadding, 16, 14, category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_FramePadding, 10, 7, category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_ItemSpacing, 10, 9, category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_ScrollbarSize, 10, category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_GrabRounding, 8, category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_FrameBorderSize, 1, category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_WindowBorderSize, 1, category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_ChildBorderSize, 1, category=dpg.mvThemeCat_Core)
        dpg.bind_theme(theme)

        with dpg.theme(tag="nav_default_theme"):
            with dpg.theme_component(dpg.mvSelectable):
                dpg.add_theme_color(dpg.mvThemeCol_Header, (255, 255, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_HeaderHovered, (248, 250, 252), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_HeaderActive, (239, 246, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Text, (51, 65, 85), category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_FrameBorderSize, 0, category=dpg.mvThemeCat_Core)

        with dpg.theme(tag="nav_active_theme"):
            with dpg.theme_component(dpg.mvSelectable):
                dpg.add_theme_color(dpg.mvThemeCol_Header, (219, 234, 254), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_HeaderHovered, (191, 219, 254), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_HeaderActive, (147, 197, 253), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Text, (30, 64, 175), category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_FrameBorderSize, 0, category=dpg.mvThemeCat_Core)

        with dpg.theme(tag="secondary_button_theme"):
            with dpg.theme_component(dpg.mvButton):
                dpg.add_theme_color(dpg.mvThemeCol_Button, (255, 255, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_ButtonHovered, (248, 250, 252), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_ButtonActive, (239, 246, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Text, (51, 65, 85), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Border, (203, 213, 225), category=dpg.mvThemeCat_Core)

        with dpg.theme(tag="primary_button_theme"):
            with dpg.theme_component(dpg.mvButton):
                dpg.add_theme_color(dpg.mvThemeCol_Button, (37, 99, 235), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_ButtonHovered, (29, 78, 216), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_ButtonActive, (30, 64, 175), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Text, (255, 255, 255), category=dpg.mvThemeCat_Core)

        with dpg.theme(tag="menu_bar_theme"):
            with dpg.theme_component(dpg.mvMenuBar):
                dpg.add_theme_color(dpg.mvThemeCol_MenuBarBg, (255, 255, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Text, (31, 41, 55), category=dpg.mvThemeCat_Core)
            with dpg.theme_component(dpg.mvMenu):
                dpg.add_theme_color(dpg.mvThemeCol_Text, (31, 41, 55), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_HeaderHovered, (239, 246, 255), category=dpg.mvThemeCat_Core)
            with dpg.theme_component(dpg.mvMenuItem):
                dpg.add_theme_color(dpg.mvThemeCol_Text, (31, 41, 55), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_HeaderHovered, (239, 246, 255), category=dpg.mvThemeCat_Core)

        with dpg.theme(tag="auth_dialog_theme"):
            with dpg.theme_component(dpg.mvWindowAppItem):
                dpg.add_theme_color(dpg.mvThemeCol_WindowBg, (255, 255, 255), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Border, (71, 85, 105), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_TitleBg, (226, 232, 240), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_TitleBgActive, (219, 234, 254), category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_WindowBorderSize, 2, category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_WindowPadding, 16, 14, category=dpg.mvThemeCat_Core)

        with dpg.theme(tag="auth_panel_theme"):
            with dpg.theme_component(dpg.mvChildWindow):
                dpg.add_theme_color(dpg.mvThemeCol_ChildBg, (248, 250, 252), category=dpg.mvThemeCat_Core)
                dpg.add_theme_color(dpg.mvThemeCol_Border, (148, 163, 184), category=dpg.mvThemeCat_Core)
                dpg.add_theme_style(dpg.mvStyleVar_ChildBorderSize, 1, category=dpg.mvThemeCat_Core)

    def _build_ui(self) -> None:
        with dpg.file_dialog(
            directory_selector=False,
            show=False,
            callback=self._on_file_selected,
            tag="file_dialog",
            width=700,
            height=420,
        ):
            dpg.add_file_extension(".pdf", color=(220, 70, 70, 255))
            dpg.add_file_extension(".pptx", color=(70, 120, 220, 255))
            dpg.add_file_extension(".lic", color=(34, 139, 94, 255))
            dpg.add_file_extension(".json", color=(34, 139, 94, 255))
            dpg.add_file_extension(".*")

        with dpg.file_dialog(
            directory_selector=True,
            show=False,
            callback=self._on_directory_selected,
            tag="directory_dialog",
            width=700,
            height=420,
        ):
            dpg.add_file_extension(".*")

        self._build_dialogs()

        with dpg.viewport_menu_bar(tag="main_menu_bar"):
            with dpg.menu(label="文件"):
                dpg.add_menu_item(label="打开上次输出", callback=self._open_last_output)
                dpg.add_menu_item(label="退出", callback=lambda *args: dpg.stop_dearpygui())
            with dpg.menu(label="授权"):
                dpg.add_menu_item(label="授权中心...", callback=self.open_auth_center)
                dpg.add_menu_item(label="查看授权状态", callback=self.show_authorization_status)
                dpg.add_separator()
                dpg.add_menu_item(label="导入离线 License...", callback=self.import_offline_license_from_menu)
                dpg.add_menu_item(label="删除现有授权", callback=self.clear_authorization)
                dpg.add_menu_item(label="导出离线设备请求...", callback=self.export_device_request_from_menu)
                dpg.add_menu_item(label="复制设备指纹", callback=self.copy_device_fingerprint)
                dpg.add_separator()
                dpg.add_menu_item(label="在线登录绑定...", callback=self.open_auth_center)
            with dpg.menu(label="帮助"):
                dpg.add_menu_item(label="关于 PDF Tools", callback=self.show_about)
            dpg.add_spacer(tag="auth_menu_spacer", width=520)
            dpg.add_text("授权：", color=(100, 116, 139))
            dpg.add_text("", tag="auth_status_badge")
        dpg.bind_item_theme("main_menu_bar", "menu_bar_theme")

        with dpg.window(tag="main_window", label=APP_TITLE, no_scrollbar=True):
            with dpg.group(horizontal=True):
                with dpg.child_window(width=220, height=-1, no_scrollbar=True):
                    dpg.add_text("功能")
                    dpg.add_separator()
                    dpg.add_listbox(
                        tag="nav_list",
                        items=[label for _, label, _ in TOOLS],
                        default_value=TOOLS[0][1],
                        num_items=len(TOOLS),
                        width=-1,
                        callback=self._on_nav_selected,
                    )

                with dpg.child_window(tag="content", width=-1, height=-1):
                    dpg.add_text("", tag="tool_title")
                    if self.title_font:
                        dpg.bind_item_font("tool_title", self.title_font)
                    dpg.add_text("", tag="tool_description", color=(100, 116, 139))
                    dpg.add_separator()
                    pdf_to_word.build_page(self)
                    pdf_merge.build_page(self)
                    pdf_split.build_page(self)
                    pdf_compress.build_page(self)
                    pdf_to_ppt_page.build_page(self)
                    pdf_to_image.build_page(self)
                    pdf_encrypt.build_page(self)
                    pdf_to_excel.build_page(self)
                    dpg.add_spacer(height=8)
                    dpg.add_separator()
                    with dpg.group(horizontal=True):
                        dpg.add_text("", tag="status_text")
                        dpg.add_loading_indicator(tag="busy_indicator", show=False, radius=2.5)
                    with dpg.group(horizontal=True):
                        dpg.add_button(label="打开", tag="open_output_btn", callback=self._open_last_output, enabled=False)
                        dpg.bind_item_theme("open_output_btn", "secondary_button_theme")
                        dpg.add_button(label="打开文件夹", tag="open_output_folder_btn", callback=self._open_output_folder, enabled=False)
                        dpg.bind_item_theme("open_output_folder_btn", "secondary_button_theme")
        self._refresh_auth_summary()
        self._show_tool("merge")
        dpg.set_frame_callback(1, self._poll_results)

    def _build_dialogs(self) -> None:
        with dpg.window(label="授权中心", tag="auth_window", show=False, width=1000, height=620, no_resize=False, no_scrollbar=True):
            dpg.add_text("授权中心", tag="auth_window_title")
            if self.title_font:
                dpg.bind_item_font("auth_window_title", self.title_font)
            with dpg.group(horizontal=True):
                dpg.add_text("当前授权", color=(100, 116, 139))
                dpg.add_text("", tag="auth_card_title")
                dpg.add_text("", tag="auth_card_detail", color=(100, 116, 139))
                dpg.add_button(label="刷新", width=70, callback=lambda *args: self._refresh_auth_summary())
                dpg.bind_item_theme(dpg.last_item(), "secondary_button_theme")
                dpg.add_button(label="删除现有授权", width=120, callback=self.clear_authorization)
                dpg.bind_item_theme(dpg.last_item(), "secondary_button_theme")

            dpg.add_separator()
            with dpg.group(horizontal=True):
                with dpg.child_window(tag="auth_side_panel", width=260, height=420, no_scrollbar=True):
                    dpg.add_text("授权方式")
                    dpg.add_radio_button(
                        items=["离线授权（推荐）", "在线登录"],
                        tag="auth_mode_selector",
                        default_value="离线授权（推荐）",
                        callback=lambda sender, app_data: self._set_auth_mode("离线授权" if app_data.startswith("离线") else "在线登录"),
                    )
                    dpg.add_separator()
                    dpg.add_text("设备信息")
                    dpg.add_text("", tag="auth_device_fingerprint", color=(100, 116, 139), wrap=230)
                    dpg.add_separator()
                    dpg.add_text("建议优先使用离线授权，适合客户内网、无法长期联网或交付后独立运行的环境。", color=(100, 116, 139), wrap=230)
                dpg.bind_item_theme("auth_side_panel", "auth_panel_theme")

                with dpg.group():
                    with dpg.child_window(tag="auth_offline_group", width=-1, height=420, no_scrollbar=True):
                        dpg.add_text("离线授权流程")
                        dpg.add_text("按下面 3 步完成后立即生效，无需重启应用。", color=(100, 116, 139))
                        dpg.add_separator()
                        self._auth_step(
                            "1",
                            "导出设备请求文件",
                            f"生成 {OFFLINE_DEVICE_REQUEST_FILENAME}，并把它带到可访问官网的电脑。",
                            "导出设备请求",
                            self.export_device_request_from_menu,
                            140,
                        )
                        self._auth_step(
                            "2",
                            "官网签发 License",
                            "在官网账号中心导入设备请求，签发并下载 .lic 文件。",
                            "打开官网离线页面",
                            self.open_offline_license_page,
                            160,
                        )
                        self._auth_step(
                            "3",
                            "导入 License 生效",
                            "回到本机导入已签发的 .lic 文件，授权状态会立即刷新。",
                            "导入已签发 License",
                            self.import_offline_license_from_menu,
                            170,
                        )
                    dpg.bind_item_theme("auth_offline_group", "auth_panel_theme")

                    with dpg.child_window(tag="auth_online_group", width=-1, height=420, no_scrollbar=True, show=False):
                        dpg.add_text("在线登录流程")
                        dpg.add_text("适合当前设备能访问官网并需要在线拉取订阅状态的场景。", color=(100, 116, 139))
                        dpg.add_separator()
                        self._auth_step("1", "生成绑定码", "本机会向服务器申请一次性设备绑定码。", "生成在线登录绑定码", self.start_online_binding, 180)
                        self._auth_step("2", "网页确认设备", "在网页端登录账号并确认绑定当前设备。", "打开网页确认", self.open_online_confirm, 130, button_tag="auth_open_confirm_btn", enabled=False)
                        dpg.add_text("3  等待自动写入订阅缓存", color=(31, 41, 55))
                        dpg.add_text("保持本窗口打开，客户端会自动轮询确认结果。", color=(100, 116, 139), wrap=620)
                        dpg.add_spacer(height=8)
                        dpg.add_text("绑定码：-", tag="auth_bind_code")
                    dpg.bind_item_theme("auth_online_group", "auth_panel_theme")

            dpg.add_text("", tag="auth_action_status", color=(100, 116, 139), wrap=900)
            dpg.add_separator()
            with dpg.group(horizontal=True):
                dpg.add_spacer(width=850)
                dpg.add_button(label="关闭", width=90, callback=lambda *args: dpg.configure_item("auth_window", show=False))
                dpg.bind_item_theme(dpg.last_item(), "secondary_button_theme")
        dpg.bind_item_theme("auth_window", "auth_dialog_theme")

        with dpg.window(label="授权状态", tag="auth_status_window", show=False, width=520, height=300):
            dpg.add_text("授权状态", tag="auth_status_title")
            if self.title_font:
                dpg.bind_item_font("auth_status_title", self.title_font)
            dpg.add_separator()
            dpg.add_text("", tag="auth_status_detail", wrap=470)
            dpg.add_spacer(height=14)
            dpg.add_button(label="关闭", width=90, callback=lambda *args: dpg.configure_item("auth_status_window", show=False))
            dpg.bind_item_theme(dpg.last_item(), "secondary_button_theme")

        with dpg.window(label="关于 PDF Tools", tag="about_window", show=False, width=480, height=230):
            dpg.add_text(APP_TITLE, tag="about_title")
            if self.title_font:
                dpg.bind_item_font("about_title", self.title_font)
            dpg.add_text("PDF 转换、合并、分割、压缩、图片导出与加密工具。", wrap=430)
            dpg.add_text("授权方式：在线登录拉取订阅，或导入离线订阅 License。", color=(100, 116, 139), wrap=430)
            dpg.add_spacer(height=14)
            dpg.add_button(label="关闭", width=90, callback=lambda *args: dpg.configure_item("about_window", show=False))
            dpg.bind_item_theme(dpg.last_item(), "secondary_button_theme")

    def _nav_button(self, label: str, tool: str) -> None:
        dpg.add_selectable(label=label, tag=f"nav_{tool}", width=-1, height=34, callback=lambda *args: self._show_tool(tool))
        dpg.bind_item_theme(f"nav_{tool}", "nav_default_theme")

    def _auth_step(
        self,
        number: str,
        title: str,
        description: str,
        button_label: str,
        callback,
        button_width: int,
        button_tag: str | None = None,
        enabled: bool = True,
    ) -> None:
        with dpg.group(horizontal=True):
            dpg.add_text(f"{number}.")
            with dpg.group():
                dpg.add_text(title)
                dpg.add_text(description, color=(100, 116, 139), wrap=600)
            if button_tag:
                dpg.add_button(label=button_label, tag=button_tag, width=button_width, callback=callback, enabled=enabled)
            else:
                dpg.add_button(label=button_label, width=button_width, callback=callback, enabled=enabled)
            dpg.bind_item_theme(dpg.last_item(), "secondary_button_theme")
        dpg.add_spacer(height=10)

    def _on_nav_selected(self, sender, app_data, user_data=None) -> None:
        del sender, user_data
        for tool, label, _description in TOOLS:
            if label == app_data:
                self._show_tool(tool)
                return

    def _file_row(self, label: str, text_tag: str, target: str) -> None:
        dpg.add_text(label)
        with dpg.group(horizontal=True):
            dpg.add_input_text(tag=text_tag, default_value="未选择文件", readonly=True, width=-120)
            dpg.add_button(label="浏览", width=96, callback=lambda *args: self._open_file_dialog(target))
            dpg.bind_item_theme(dpg.last_item(), "secondary_button_theme")

    def _dir_row(self, label: str, value_tag: str) -> None:
        dpg.add_text(label)
        with dpg.group(horizontal=True):
            dpg.add_input_text(tag=value_tag, width=-120)
            dpg.add_button(label="浏览", width=96, callback=lambda *args: self._open_directory_dialog(value_tag))
            dpg.bind_item_theme(dpg.last_item(), "secondary_button_theme")

    def page_header(self, title: str, description: str) -> None:
        del title, description
        dpg.add_spacer(height=2)

    def primary_button(self, label: str, callback: Callable[..., None], tag: str | None = None, width: int = 130) -> None:
        if tag:
            dpg.add_button(label=label, tag=tag, width=width, height=34, callback=callback)
        else:
            dpg.add_button(label=label, width=width, height=34, callback=callback)
        dpg.bind_item_theme(dpg.last_item(), "primary_button_theme")

    def secondary_button(self, label: str, callback: Callable[..., None], width: int = 120) -> None:
        dpg.add_button(label=label, width=width, height=34, callback=callback)
        dpg.bind_item_theme(dpg.last_item(), "secondary_button_theme")

    def file_row(self, label: str, text_tag: str, target: str) -> None:
        self._file_row(label, text_tag, target)

    def dir_row(self, label: str, value_tag: str) -> None:
        self._dir_row(label, value_tag)

    def open_file_dialog(self, target: str) -> None:
        self._open_file_dialog(target)

    def merge_remove(self, *args) -> None:
        self._merge_remove(*args)

    def merge_move(self, delta: int) -> None:
        self._merge_move(delta)

    def run_word(self, *args) -> None:
        self._run_word(*args)

    def run_merge(self, *args) -> None:
        self._run_merge(*args)

    def run_split(self, *args) -> None:
        self._run_split(*args)

    def run_compress(self, *args) -> None:
        self._run_compress(*args)

    def run_ppt(self, *args) -> None:
        self._run_ppt(*args)

    def run_image(self, *args) -> None:
        self._run_image(*args)

    def run_encrypt(self, *args) -> None:
        self._run_encrypt(*args)

    def _show_tool(self, tool: str) -> None:
        self.current_tool = tool
        for item, label, description in TOOLS:
            dpg.configure_item(f"tool_{item}", show=item == tool)
            if item == tool:
                if dpg.does_item_exist("nav_list"):
                    dpg.set_value("nav_list", label)
                dpg.set_value("tool_title", label)
                dpg.set_value("tool_description", description)
        self._set_status("")

    def _open_file_dialog(self, target: str) -> None:
        paths = self._open_native_file_dialog(target)
        if paths is not NATIVE_DIALOG_UNAVAILABLE:
            if not paths:
                return
            if target == "auth_license":
                self._import_offline_license(paths[0])
            elif target == "merge_add":
                self._add_merge_files(paths)
            else:
                self._set_file_target(target, paths[0])
            return
        self._file_dialog_target = target
        dpg.show_item("file_dialog")

    def _open_directory_dialog(self, target: str) -> None:
        path = self._open_native_directory_dialog()
        if path is not NATIVE_DIALOG_UNAVAILABLE:
            if path is None:
                return
            if target == "auth_export_dir":
                self._export_device_request(path / OFFLINE_DEVICE_REQUEST_FILENAME)
            else:
                dpg.set_value(target, str(path))
            return
        self._file_dialog_target = target
        dpg.show_item("directory_dialog")

    def _open_native_file_dialog(self, target: str) -> list[Path] | object:
        try:
            from tkinter import Tk, filedialog

            root = Tk()
            root.withdraw()
            root.attributes("-topmost", True)
            if target == "auth_license":
                filetypes = [("License files", "*.lic *.json"), ("All files", "*.*")]
            elif target == "ppt_output":
                filetypes = [("PowerPoint files", "*.pptx"), ("All files", "*.*")]
            else:
                filetypes = [("PDF files", "*.pdf"), ("All files", "*.*")]
            if target == "merge_add":
                value = filedialog.askopenfilenames(parent=root, filetypes=filetypes)
            else:
                single = filedialog.askopenfilename(parent=root, filetypes=filetypes)
                value = (single,) if single else ()
            root.destroy()
        except Exception:
            return NATIVE_DIALOG_UNAVAILABLE
        return [Path(item) for item in value] if value else []

    def _open_native_directory_dialog(self) -> Path | object | None:
        try:
            from tkinter import Tk, filedialog

            root = Tk()
            root.withdraw()
            root.attributes("-topmost", True)
            value = filedialog.askdirectory(parent=root)
            root.destroy()
        except Exception:
            return NATIVE_DIALOG_UNAVAILABLE
        return Path(value) if value else None

    def _on_file_selected(self, sender, app_data, user_data=None) -> None:
        del sender
        target = self._file_dialog_target
        if not target:
            return
        selections = app_data.get("selections") or {}
        if not selections:
            return
        paths = [Path(path) for path in selections.values()]
        if target == "auth_license":
            self._import_offline_license(paths[0])
            return
        if target == "merge_add":
            self._add_merge_files(paths)
            return
        self._set_file_target(target, paths[0])

    def _on_directory_selected(self, sender, app_data, user_data=None) -> None:
        del sender
        target = self._file_dialog_target
        if not target:
            return
        path = app_data.get("file_path_name")
        if path and target == "auth_export_dir":
            self._export_device_request(Path(path) / OFFLINE_DEVICE_REQUEST_FILENAME)
        elif path:
            dpg.set_value(target, path)

    def _set_file_target(self, target: str, path: Path) -> None:
        if target == "merge_add":
            self._add_merge_files([path])
            return

        setattr(self, target, str(path))
        dpg.set_value(f"{target}_text", str(path))
        if target == "word_pdf":
            if not dpg.get_value("word_out_dir"):
                dpg.set_value("word_out_dir", str(path.parent))
            dpg.set_value("word_file", path.with_suffix(".docx").name)
        elif target == "split_pdf":
            if not dpg.get_value("split_out_dir"):
                dpg.set_value("split_out_dir", str(path.parent))
            if not dpg.get_value("split_prefix"):
                dpg.set_value("split_prefix", path.stem)
        elif target == "compress_pdf" and not dpg.get_value("compress_output"):
            dpg.set_value("compress_output", str(path.with_name(path.stem + "_compressed.pdf")))
        elif target == "ppt_pdf" and not dpg.get_value("ppt_output"):
            dpg.set_value("ppt_output", str(path.with_suffix(".pptx")))
        elif target == "image_pdf":
            if not dpg.get_value("image_out_dir"):
                dpg.set_value("image_out_dir", str(path.parent))
            if not dpg.get_value("image_single_name"):
                dpg.set_value("image_single_name", path.stem + "." + dpg.get_value("image_format"))
        elif target == "encrypt_pdf" and not dpg.get_value("encrypt_output"):
            dpg.set_value("encrypt_output", str(path.with_name(path.stem + "_encrypted.pdf")))

    def _add_merge_files(self, paths: list[Path]) -> None:
        if not paths:
            return
        self.merge_files.extend(str(path) for path in paths)
        dpg.configure_item("merge_list", items=[Path(path).name for path in self.merge_files])
        if not dpg.get_value("merge_output"):
            first = paths[0]
            dpg.set_value("merge_output", str(first.with_name(first.stem + "_merged.pdf")))

    def _merge_remove(self, *args) -> None:
        selected = dpg.get_value("merge_list")
        if not selected:
            return
        names = [Path(p).name for p in self.merge_files]
        if selected in names:
            self.merge_files.pop(names.index(selected))
            dpg.configure_item("merge_list", items=[Path(p).name for p in self.merge_files])

    def _merge_move(self, delta: int) -> None:
        selected = dpg.get_value("merge_list")
        names = [Path(p).name for p in self.merge_files]
        if selected not in names:
            return
        idx = names.index(selected)
        new_idx = idx + delta
        if not 0 <= new_idx < len(self.merge_files):
            return
        self.merge_files[idx], self.merge_files[new_idx] = self.merge_files[new_idx], self.merge_files[idx]
        dpg.configure_item("merge_list", items=[Path(p).name for p in self.merge_files])
        dpg.set_value("merge_list", Path(self.merge_files[new_idx]).name)

    def _run_word(self, *args) -> None:
        inp = self._require_file(self.word_pdf, "Select an input PDF.")
        if not inp:
            return
        out_dir = dpg.get_value("word_out_dir").strip()
        name = dpg.get_value("word_file").strip()
        if not out_dir or not name:
            self._error("Choose an output folder and file name.")
            return
        out = Path(out_dir) / name
        if out.exists() and not dpg.get_value("word_overwrite"):
            self._error("Output exists. Enable overwrite or choose another file.")
            return
        start = dpg.get_value("word_start") or 0
        end = dpg.get_value("word_end") or 0
        self._run_task(
            lambda: convert_pdf_to_docx(str(inp), str(out), start=start if start else None, end=end if end else None),
            str(out),
            "正在转换 PDF 为 Word...",
            feature_key="word",
        )

    def _run_merge(self, *args) -> None:
        if len(self.merge_files) < 2:
            self._error("Add at least two PDF files.")
            return
        out = dpg.get_value("merge_output").strip()
        if not out:
            self._error("Choose an output PDF.")
            return
        self._run_task(lambda: merge_pdfs(self.merge_files, out), out, "正在合并 PDF...", feature_key="merge")

    def _run_split(self, *args) -> None:
        inp = self._require_file(self.split_pdf, "Select an input PDF.")
        if not inp:
            return
        out_dir = dpg.get_value("split_out_dir").strip()
        if not out_dir:
            self._error("Choose an output folder.")
            return
        prefix = dpg.get_value("split_prefix").strip() or inp.stem
        pages = dpg.get_value("split_pages")
        self._run_task(lambda: split_pdf(str(inp), out_dir, pages, prefix), out_dir, "正在分割 PDF...", feature_key="split")

    def _run_compress(self, *args) -> None:
        inp = self._require_file(self.compress_pdf, "Select an input PDF.")
        out = dpg.get_value("compress_output").strip()
        if not inp or not out:
            self._error("Select an input PDF and output PDF.")
            return
        self._run_task(lambda: compress_pdf(str(inp), out), out, "正在压缩 PDF...", feature_key="compress")

    def _run_ppt(self, *args) -> None:
        inp = self._require_file(self.ppt_pdf, "Select an input PDF.")
        out = dpg.get_value("ppt_output").strip()
        if not inp or not out:
            self._error("Select an input PDF and output PPTX.")
            return
        if not out.lower().endswith(".pptx"):
            out += ".pptx"
        dpi = dpg.get_value("ppt_dpi")
        self._run_task(lambda: pdf_to_ppt(str(inp), out, dpi), out, "正在转换 PDF 为 PPT...", feature_key="ppt")

    def _run_image(self, *args) -> None:
        inp = self._require_file(self.image_pdf, "Select an input PDF.")
        out_dir = dpg.get_value("image_out_dir").strip()
        if not inp or not out_dir:
            self._error("Select an input PDF and output folder.")
            return
        fmt = dpg.get_value("image_format")
        dpi = dpg.get_value("image_dpi")
        mode = "pages" if dpg.get_value("image_mode") in {"One file per page", "逐页导出"} else "single"
        single_name = dpg.get_value("image_single_name").strip()

        def task() -> str:
            return pdf_to_images(str(inp), out_dir, fmt, dpi, mode, single_name)

        self._run_task(task, None, "正在转换 PDF 为图片...", feature_key="image")

    def _run_encrypt(self, *args) -> None:
        inp = self._require_file(self.encrypt_pdf, "Select an input PDF.")
        out = dpg.get_value("encrypt_output").strip()
        password = dpg.get_value("encrypt_password")
        if not inp or not out:
            self._error("Select an input PDF and output PDF.")
            return
        if not password:
            self._error("Enter a password.")
            return
        self._run_task(lambda: encrypt_pdf(str(inp), out, password), out, "正在加密 PDF...", feature_key="encrypt")

    def open_auth_center(self, *args) -> None:
        self._refresh_auth_summary()
        self._set_auth_mode("离线授权")
        self._show_window_centered("auth_window", 1000, 620)

    def _set_auth_mode(self, mode: str) -> None:
        if dpg.does_item_exist("auth_offline_group"):
            dpg.configure_item("auth_offline_group", show=mode == "离线授权")
        if dpg.does_item_exist("auth_online_group"):
            dpg.configure_item("auth_online_group", show=mode == "在线登录")
        if dpg.does_item_exist("auth_mode_selector"):
            dpg.set_value("auth_mode_selector", "离线授权（推荐）" if mode == "离线授权" else "在线登录")

    def show_authorization_status(self, *args) -> None:
        summary = self._refresh_auth_summary()
        dpg.set_value(
            "auth_status_detail",
            "\n".join(
                [
                    f"状态：{summary['status']}",
                    f"方式：{summary['source']}",
                    f"有效期至：{summary['expires_at']}",
                    f"剩余时间：{summary['remaining']}",
                    f"授权版本：{summary['edition']}",
                    f"说明：{summary['message']}",
                ]
            ),
        )
        self._show_window_centered("auth_status_window", 520, 300)

    def import_offline_license_from_menu(self, *args) -> None:
        self._open_file_dialog("auth_license")

    def export_device_request_from_menu(self, *args) -> None:
        self._open_directory_dialog("auth_export_dir")

    def copy_device_fingerprint(self, *args) -> None:
        dpg.set_clipboard_text(get_device_fingerprint())
        self._set_auth_action_status("设备指纹已复制。")
        self._set_status("设备指纹已复制。")

    def clear_authorization(self, *args) -> None:
        self.auth_store.clear()
        self.poll_code = None
        self.confirm_url = None
        self._refresh_auth_summary()
        self._set_auth_action_status("现有授权已删除。")
        self._set_status("现有授权已删除。")

    def open_offline_license_page(self, *args) -> None:
        webbrowser.open(f"{DEFAULT_SERVER_URL}/account/licenses/offline")

    def start_online_binding(self, *args) -> None:
        self._set_auth_action_status("正在生成绑定码...")

        def worker() -> None:
            try:
                data = start_device_binding()
                self._results.put(("auth_bind", data))
            except Exception as exc:
                self._results.put(("auth_error", str(exc)))

        threading.Thread(target=worker, daemon=True).start()

    def open_online_confirm(self, *args) -> None:
        if self.confirm_url:
            webbrowser.open(f"{DEFAULT_SERVER_URL}{self.confirm_url}")

    def show_about(self, *args) -> None:
        self._show_window_centered("about_window", 480, 230)

    def _show_window_centered(self, tag: str, width: int, height: int) -> None:
        viewport_width = dpg.get_viewport_client_width() or VIEWPORT_WIDTH
        viewport_height = dpg.get_viewport_client_height() or VIEWPORT_HEIGHT
        x = max(16, int((viewport_width - width) / 2))
        y = max(48, int((viewport_height - height) / 2))
        dpg.configure_item(tag, pos=(x, y), show=True)

    def _refresh_auth_summary(self) -> dict[str, str]:
        result = get_authorization_status(self.auth_store)
        summary = get_authorization_summary(result)
        badge = f"{summary['status']} · {summary['source']}"
        detail = (
            f"有效期：{summary['expires_at']}    "
            f"剩余：{summary['remaining']}    "
            f"版本：{summary['edition']}"
        )
        status_detail = "\n".join(
            [
                f"有效期至：{summary['expires_at']}",
                f"剩余时间：{summary['remaining']}",
                f"授权版本：{summary['edition']}",
                f"状态说明：{summary['message']}",
            ]
        )
        for tag, value in {
            "auth_status_badge": badge,
            "auth_card_title": badge,
            "auth_card_detail": detail,
            "auth_device_fingerprint": f"设备指纹：{get_device_fingerprint()[:28]}...",
        }.items():
            if dpg.does_item_exist(tag):
                dpg.set_value(tag, value)
        if dpg.does_item_exist("auth_status_detail"):
            dpg.set_value("auth_status_detail", status_detail)
        return summary

    def _set_auth_action_status(self, message: str) -> None:
        if dpg.does_item_exist("auth_action_status"):
            dpg.set_value("auth_action_status", message)

    def _import_offline_license(self, path: Path) -> None:
        result = import_offline_license(path, self.auth_store)
        self._refresh_auth_summary()
        message = result.message if result.valid else f"导入失败：{result.message}"
        self._set_auth_action_status(message)
        self._set_status(message)

    def _export_device_request(self, path: Path) -> None:
        try:
            save_offline_device_request(path)
        except OSError as exc:
            message = f"导出失败：{exc}"
        else:
            message = f"离线设备请求已保存：{path}"
            open_parent_folder(str(path))
        self._set_auth_action_status(message)
        self._set_status(message)

    def _handle_auth_bind_result(self, data: dict[str, Any]) -> None:
        self.poll_code = str(data["pollCode"])
        self.confirm_url = str(data["confirmUrl"])
        dpg.set_value("auth_bind_code", f"绑定码：{data['userCode']}")
        dpg.configure_item("auth_open_confirm_btn", enabled=True)
        self._set_auth_action_status("绑定码已生成，请在网页登录并确认本机设备。")

    def _poll_authorization_binding(self) -> None:
        if not self.poll_code:
            return
        now = time.monotonic()
        if now - self.last_poll_at < 3:
            return
        self.last_poll_at = now

        def worker() -> None:
            try:
                data = poll_device_binding(self.poll_code or "", self.auth_store)
                self._results.put(("auth_poll", data))
            except Exception as exc:
                self._results.put(("auth_error", str(exc)))

        threading.Thread(target=worker, daemon=True).start()

    def _require_file(self, value: str | None, message: str) -> Path | None:
        if not value:
            self._error(message)
            return None
        path = Path(value)
        if not path.is_file():
            self._error("Input file does not exist.")
            return None
        return path

    def _run_task(
        self,
        func: Callable[[], str | None],
        expected_output: str | None,
        message: str,
        feature_key: str | None = None,
    ) -> None:
        if self.busy:
            self._set_status("当前任务仍在运行，请等待完成。", busy=True)
            return
        if feature_key and not self._allow_feature_run(feature_key):
            return
        self.busy = True
        self._set_task_buttons_enabled(False)
        self._set_status(message, busy=True)

        def worker() -> None:
            try:
                result = func()
                out = result or expected_output
                self._results.put(("done", out))
            except Exception as exc:
                self._results.put(("error", str(exc)))

        threading.Thread(target=worker, daemon=True).start()

    def _allow_feature_run(self, feature_key: str) -> bool:
        if get_authorization_status(self.auth_store).valid:
            return True
        if self.trial_usage_store.try_consume(feature_key):
            return True

        feature_label = TOOL_LABELS.get(feature_key, feature_key)
        message = f"未授权版本每天每个功能只能执行 {TRIAL_DAILY_LIMIT} 次，{feature_label} 今日次数已用完。"
        self._set_auth_action_status(message)
        self._error(message)
        return False

    def _poll_results(self, *args) -> None:
        self._update_menu_status_layout()
        self._poll_authorization_binding()
        while not self._results.empty():
            status, payload = self._results.get_nowait()
            if status == "done":
                self._task_done(payload)
            elif status == "error":
                self._task_failed(payload or "Unknown error")
            elif status == "auth_bind" and isinstance(payload, dict):
                self._handle_auth_bind_result(payload)
            elif status == "auth_poll" and isinstance(payload, dict):
                self._handle_auth_poll_result(payload)
            elif status == "auth_error":
                self._set_auth_action_status(str(payload))
        dpg.set_frame_callback(dpg.get_frame_count() + 1, self._poll_results)

    def _update_menu_status_layout(self) -> None:
        if not dpg.does_item_exist("auth_menu_spacer"):
            return
        try:
            width = dpg.get_viewport_client_width() or VIEWPORT_WIDTH
        except Exception:
            width = VIEWPORT_WIDTH
        dpg.configure_item("auth_menu_spacer", width=max(24, width - 350))

    def _handle_auth_poll_result(self, data: dict[str, Any]) -> None:
        if data.get("status") == "confirmed":
            self.poll_code = None
            self._refresh_auth_summary()
            self._set_auth_action_status("在线登录成功，已拉取订阅状态。")
        elif data.get("status") in {"expired", "revoked"}:
            self.poll_code = None
            self._set_auth_action_status("绑定码已过期或已撤销，请重新生成。")

    def _task_done(self, output: str | None) -> None:
        self.busy = False
        self._set_task_buttons_enabled(True)
        self.output_path = output
        dpg.configure_item("open_output_btn", enabled=bool(output))
        dpg.configure_item("open_output_folder_btn", enabled=bool(output))
        self._set_status(f"完成：{output}" if output else "完成。", busy=False)

    def _task_failed(self, message: str) -> None:
        self.busy = False
        self._set_task_buttons_enabled(True)
        self._error(message)
        self._set_status("失败。", busy=False)

    def _set_task_buttons_enabled(self, enabled: bool) -> None:
        for tag in TASK_BUTTON_TAGS:
            if dpg.does_item_exist(tag):
                dpg.configure_item(tag, enabled=enabled)

    def _set_status(self, message: str, busy: bool = False) -> None:
        dpg.set_value("status_text", message)
        dpg.configure_item("busy_indicator", show=busy)

    def _error(self, message: str) -> None:
        dpg.set_value("status_text", f"错误：{message}")

    def _open_last_output(self, *args) -> None:
        if self.output_path:
            open_path(self.output_path)

    def _open_output_folder(self, *args) -> None:
        if self.output_path:
            open_parent_folder(self.output_path)


def merge_pdfs(inputs: list[str], output: str) -> str:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for pdf in inputs:
        reader = PdfReader(pdf)
        for page in reader.pages:
            writer.add_page(page)
    out = Path(output)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("wb") as file:
        writer.write(file)
    return str(out)


def split_pdf(pdf_path: str, out_dir: str, pages_per_part: int, base_name: str) -> str:
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(pdf_path)
    out_path = Path(out_dir)
    out_path.mkdir(parents=True, exist_ok=True)
    part = 1
    for start in range(0, len(reader.pages), pages_per_part):
        writer = PdfWriter()
        for index in range(start, min(start + pages_per_part, len(reader.pages))):
            writer.add_page(reader.pages[index])
        out = out_path / f"{base_name}_part{part:03d}.pdf"
        with out.open("wb") as file:
            writer.write(file)
        part += 1
    return str(out_path)


def compress_pdf(pdf_path: str, out_path: str) -> str:
    from pypdf import PdfWriter

    writer = PdfWriter(clone_from=pdf_path)
    for page in writer.pages:
        page.compress_content_streams(level=9)
    writer.compress_identical_objects(remove_identicals=True, remove_orphans=True)
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("wb") as file:
        writer.write(file)
    return str(out)


def encrypt_pdf(pdf_path: str, out_path: str, password: str) -> str:
    from pypdf import PdfReader, PdfWriter

    reader = PdfReader(pdf_path)
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(password, algorithm="AES-256")
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("wb") as file:
        writer.write(file)
    return str(out)


def pdf_to_ppt(pdf_path: str, out_path: str, dpi: int) -> str:
    import fitz
    from pptx import Presentation

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(pdf_path)
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    slide_w = prs.slide_width
    slide_h = prs.slide_height

    with TemporaryDirectory() as td:
        td_path = Path(td)
        for index in range(doc.page_count):
            page = doc.load_page(index)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            png = td_path / f"page_{index + 1:03d}.png"
            pix.save(str(png))
            slide = prs.slides.add_slide(blank_layout)
            pic = slide.shapes.add_picture(str(png), 0, 0)
            scale = min(slide_w / pic.width, slide_h / pic.height)
            pic.width = int(pic.width * scale)
            pic.height = int(pic.height * scale)
            pic.left = int((slide_w - pic.width) / 2)
            pic.top = int((slide_h - pic.height) / 2)

    prs.save(str(out))
    return str(out)


def pdf_to_images(pdf_path: str, out_dir: str, fmt: str, dpi: int, mode: str, single_name: str) -> str:
    import fitz
    from PIL import Image

    out_path = Path(out_dir)
    out_path.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(pdf_path)
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    fmt = fmt.lower()

    if mode == "pages":
        for index in range(doc.page_count):
            page = doc.load_page(index)
            pix = page.get_pixmap(matrix=mat, alpha=False)
            out = out_path / f"{Path(pdf_path).stem}_page{index + 1:03d}.{fmt}"
            if fmt == "png":
                pix.save(str(out))
            else:
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                if image.mode != "RGB":
                    image = image.convert("RGB")
                image.save(out, format="JPEG" if fmt in {"jpg", "jpeg"} else fmt.upper(), quality=95)
        return str(out_path)

    images: list[Image.Image] = []
    for index in range(doc.page_count):
        page = doc.load_page(index)
        pix = page.get_pixmap(matrix=mat, alpha=False)
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        if image.mode != "RGB":
            image = image.convert("RGB")
        images.append(image)
    if not images:
        raise RuntimeError("PDF has no pages to convert.")

    width = max(image.width for image in images)
    height = sum(image.height for image in images)
    canvas = Image.new("RGB", (width, height), (255, 255, 255))
    y = 0
    for image in images:
        canvas.paste(image, ((width - image.width) // 2, y))
        y += image.height

    name = single_name.strip() or f"{Path(pdf_path).stem}.{fmt}"
    if not name.lower().endswith(f".{fmt}"):
        name += f".{fmt}"
    out = out_path / name
    canvas.save(out, format="JPEG" if fmt in {"jpg", "jpeg"} else fmt.upper(), quality=95)
    return str(out)


def open_path(path: str) -> None:
    target = Path(path)
    if sys.platform.startswith("linux"):
        subprocess.Popen(["xdg-open", str(target)])
    else:
        webbrowser.open(target.as_uri() if target.exists() else path)


def open_parent_folder(path: str) -> None:
    target = Path(path)
    if sys.platform.startswith("win"):
        if target.exists():
            subprocess.Popen(["explorer", "/select,", str(target)])
        else:
            subprocess.Popen(["explorer", str(target.parent)])
    elif sys.platform.startswith("linux"):
        subprocess.Popen(["xdg-open", str(target.parent)])
    else:
        webbrowser.open(target.parent.as_uri())


def main() -> int:
    return PdfToolsApp().run()
